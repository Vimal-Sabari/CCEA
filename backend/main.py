"""
FastAPI wrapper around ccea_core.

Run locally:
    cd cce-app/backend
    python -m venv .venv && .venv\\Scripts\\activate      (Windows)
    pip install -r requirements.txt
    uvicorn main:app --reload --port 8000
"""

from __future__ import annotations

import csv
import io

from fastapi import FastAPI, File, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse

import ccea_core
import llm

app = FastAPI(title="CCEA Optimizer", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],            # local dev; tighten before deploy
    allow_methods=["*"],
    allow_headers=["*"],
)


# ----------------------------------------------------------------------
# file parsing
# ----------------------------------------------------------------------
def _rows_from_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(r) for r in reader]


def _rows_from_xlsx(data: bytes) -> list[dict]:
    """Flatten the multi-sheet CCE_AllSamples.xlsx (one sheet per sample,
    plus optional _README / _AttackScore_Derivation sheets which are skipped)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows: list[dict] = []
    for sheet in wb.sheetnames:
        if sheet.startswith("_"):
            continue
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else "" for h in next(it)]
        except StopIteration:
            continue
        if "Candidate_ID" not in header or "Signature_ID" not in header:
            continue
        for raw in it:
            if raw is None or all(v is None for v in raw):
                continue
            rec = {header[i]: raw[i] for i in range(min(len(header), len(raw)))}
            if rec.get("Candidate_ID") in (None, ""):
                continue
            rows.append(rec)
    return rows


def _parse_upload(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _rows_from_xlsx(data)
    return _rows_from_csv(data)


# ----------------------------------------------------------------------
# routes
# ----------------------------------------------------------------------
@app.get("/api/health")
def health() -> dict:
    return {
        "status": "ok",
        "parameters": ccea_core.PARAMETERS,
        "reduced_severity": llm.status(),
    }


@app.post("/api/analyze")
async def analyze(
    file: UploadFile = File(...),
    download: str | None = Query(None, description="pass 'csv' to get the optimal combinations as a CSV"),
    use_llm: bool | None = Query(
        None,
        description="true = force the LLM estimator, false = force the CSV placeholder, "
                    "omit = follow the LLM_ENABLED env var",
    ),
    generations: int | None = Query(
        None,
        description="number of generations (defaults to MAX_GENERATIONS=50; all generations always run when LLM is active)",
    ),
):
    data = await file.read()
    try:
        rows = _parse_upload(file.filename, data)
        estimator = llm.get_estimator(force=use_llm)
        max_gens = generations if generations is not None else ccea_core.MAX_GENERATIONS
        # When LLM is active, run all generations (no early stopping) so the
        # full set of generations requested by the user is always executed.
        patience = max_gens if estimator is not None else None
        result = ccea_core.run_all(rows, reduced_severity_fn=estimator, max_generations=max_gens, patience=patience)
    except ccea_core.ValidationError as exc:
        return JSONResponse(status_code=422, content={"error": str(exc)})
    except Exception as exc:  # noqa: BLE001 - surface any parse/compute failure to the client
        return JSONResponse(status_code=400, content={"error": f"{type(exc).__name__}: {exc}"})

    if download == "csv":
        return _combinations_csv(result)
    return result


def _combinations_csv(result: dict) -> StreamingResponse:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([
        "Sample", "Hash", "Best_Fitness", "Best_Generation",
        "Candidate_ID", "Signature_ID", "Defense_ID",
        "Attack_Plan_Score", "Attack_Mapped", "Reduced_Severity", "Reduced_Severity_Mapped", "Fitness",
    ])
    for s in result["samples"]:
        for c in s["optimal_combination"]:
            w.writerow([
                s["label"], s["hash"], s["best_fitness"], s["best_generation"],
                c["candidate_id"], c["signature_id"], c["defense_id"],
                c["attack_plan_score"], c["attack_mapped"], c["reduced_severity"], c["reduced_severity_mapped"],
                c.get("fitness", 0.0),
            ])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=optimal_defense_combinations.csv"},
    )
